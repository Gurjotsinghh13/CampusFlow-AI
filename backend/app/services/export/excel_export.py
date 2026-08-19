from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.utils.timeline import PeriodSlot


def build_excel(
    title: str,
    days: list[str],
    periods_per_day: int,
    period_slots: list[PeriodSlot],
    grid: dict[str, list[str]],
) -> bytes:
    wb = Workbook()
    ws = wb.active
    if ws is None:
        raise RuntimeError("Unable to create Excel worksheet")
    ws.title = "Timetable"

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=10)
    border_side = Side(style="thin", color="CBD5E1")
    thin_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    wrap_center = Alignment(wrap_text=True, horizontal="center", vertical="center")

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=periods_per_day + 1)
    title_cell = ws.cell(row=1, column=1, value=title)
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="center")

    header_row = 2
    ws.cell(row=header_row, column=1, value="Day / Period").font = header_font
    ws.cell(row=header_row, column=1).fill = header_fill
    ws.cell(row=header_row, column=1).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for period in range(periods_per_day):
        if period < len(period_slots):
            slot = period_slots[period]
            header_val = f"{slot.label}\n({slot.start_time}–{slot.end_time})"
        else:
            header_val = f"P{period + 1}"

        cell = ws.cell(row=header_row, column=period + 2, value=header_val)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_offset, day in enumerate(days):
        row = header_row + 1 + row_offset
        day_cell = ws.cell(row=row, column=1, value=day.title())
        day_cell.font = Font(bold=True)
        day_cell.border = thin_border
        day_cell.alignment = Alignment(horizontal="center", vertical="center")

        for period in range(periods_per_day):
            cell = ws.cell(row=row, column=period + 2, value=grid[day][period] or "")
            cell.alignment = wrap_center
            cell.border = thin_border

    ws.column_dimensions["A"].width = 16
    for period in range(periods_per_day):
        ws.column_dimensions[get_column_letter(period + 2)].width = 22

    ws.row_dimensions[header_row].height = 28
    for row_offset in range(len(days)):
        ws.row_dimensions[header_row + 1 + row_offset].height = 54

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
